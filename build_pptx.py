#!/usr/bin/env python3
"""Build gdr_gds.pptx from hardcoded content extracted from gdr_gds.md."""

import os
import textwrap
import io

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

import cairosvg

SRC_DIR = os.path.dirname(os.path.abspath(__file__))
DST = os.path.join(SRC_DIR, "gdr_gds.pptx")
FIGS = os.path.join(SRC_DIR, "figs")

# ── color palette ──────────────────────────────────────────────────
GREEN_DARK = RGBColor(0x1B, 0x5E, 0x20)
GREEN_LIGHT = RGBColor(0x4C, 0xAF, 0x50)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
LIGHT_GRAY = RGBColor(0xF0, 0xF2, 0xF5)
CODE_BG = RGBColor(0x1E, 0x1E, 0x2E)
CODE_FG = RGBColor(0xE0, 0xE0, 0xE0)
TABLE_BORDER = RGBColor(0xB0, 0xB0, 0xB0)
TABLE_HDR_BG = GREEN_DARK
TABLE_HDR_FG = WHITE
TABLE_STRIPE = RGBColor(0xF7, 0xF8, 0xFA)

# slide dimensions (16:9)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

TITLE_TOP = Inches(0)
TITLE_H = Inches(1.0)
BODY_TOP = Inches(1.3)
FOOTER_TOP = Inches(7.0)
FOOTER_H = Inches(0.5)


def blank_slide(pres):
    slide = pres.slides.add_slide(pres.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    return slide


def add_rect(slide, left, top, width, height, fill_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    return shape


def add_textbox(slide, left, top, width, height):
    txbox = slide.shapes.add_textbox(left, top, width, height)
    txbox.text_frame.word_wrap = True
    return txbox


def set_para(para, text, font_size=Pt(13), color=BLACK, bold=False, alignment=None,
             font_name=None, space_after=None, space_before=None):
    para.text = text
    for run in para.runs:
        run.font.size = font_size
        run.font.color.rgb = color
        run.font.bold = bold
        if font_name:
            run.font.name = font_name
    if alignment is not None:
        para.alignment = alignment
    if space_after is not None:
        para.space_after = space_after
    if space_before is not None:
        para.space_before = space_before
    return para


def add_para(tf, text, font_size=Pt(13), color=BLACK, bold=False, alignment=None,
             font_name=None, space_after=Pt(4), space_before=None, level=0):
    p = tf.add_paragraph()
    p.level = level
    return set_para(p, text, font_size=font_size, color=color, bold=bold,
                    alignment=alignment, font_name=font_name,
                    space_after=space_after, space_before=space_before)


def title_bar(slide, title_text, subtitle_text=None):
    add_rect(slide, Inches(0), TITLE_TOP, SLIDE_W, TITLE_H, GREEN_DARK)
    tb = add_textbox(slide, Inches(0.8), Inches(0.18), Inches(11.7), Inches(0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    set_para(tf.paragraphs[0], title_text, Pt(22), WHITE, bold=True,
             alignment=PP_ALIGN.LEFT)
    if subtitle_text:
        add_para(tf, subtitle_text, Pt(13), RGBColor(0xC8, 0xE6, 0xC9), bold=False,
                 alignment=PP_ALIGN.LEFT)


def add_footer(slide, page_num):
    add_rect(slide, Inches(0), FOOTER_TOP, SLIDE_W, Pt(1.5), GREEN_LIGHT)
    tb = add_textbox(slide, Inches(0.4), Inches(7.05), Inches(5), Inches(0.4))
    tf = tb.text_frame
    set_para(tf.paragraphs[0], "NVIDIA GPUDirect", Pt(9), DARK_GRAY, alignment=PP_ALIGN.LEFT)
    tb2 = add_textbox(slide, Inches(11.5), Inches(7.05), Inches(1.5), Inches(0.4))
    tf2 = tb2.text_frame
    set_para(tf2.paragraphs[0], str(page_num), Pt(9), DARK_GRAY, alignment=PP_ALIGN.RIGHT)


def bullet_slide(pres, title, items, subtitle=None):
    slide = blank_slide(pres)
    title_bar(slide, title, subtitle)
    tb = add_textbox(slide, Inches(0.8), BODY_TOP, Inches(11.7), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True
    set_para(tf.paragraphs[0], items[0], Pt(14), BLACK, bold=False,
             alignment=PP_ALIGN.LEFT)
    for item in items[1:]:
        add_para(tf, item, Pt(13), DARK_GRAY, space_after=Pt(6))
    add_footer(slide, pres.slides.index(slide) + 1)
    return slide


def table_slide(pres, title, headers, rows, subtitle=None, col_widths=None, notes=None):
    slide = blank_slide(pres)
    title_bar(slide, title, subtitle)
    nrows = len(rows) + 1
    ncols = len(headers)
    tbl_left = Inches(0.8)
    tbl_top = BODY_TOP
    tbl_w = Inches(11.7)
    tbl_h = Inches(min(0.45 * nrows, 5.2))

    shape = slide.shapes.add_table(nrows, ncols, tbl_left, tbl_top, tbl_w, tbl_h)
    table = shape.table

    if col_widths:
        for ci, w in enumerate(col_widths):
            table.columns[ci].width = w

    for ci, hdr in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = hdr
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(12)
            p.font.color.rgb = WHITE
            p.font.bold = True
            p.alignment = PP_ALIGN.LEFT
        tcPr = cell._tc.get_or_add_tcPr()
        solidFill = cell._tc.makeelement(qn('a:solidFill'), {})
        srgbClr = cell._tc.makeelement(qn('a:srgbClr'), {'val': '1B5E20'})
        solidFill.append(srgbClr)
        tcPr.append(solidFill)

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri + 1, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = BLACK
                p.alignment = PP_ALIGN.LEFT
            if ri % 2 == 1:
                tcPr = cell._tc.get_or_add_tcPr()
                sf = cell._tc.makeelement(qn('a:solidFill'), {})
                sc = cell._tc.makeelement(qn('a:srgbClr'), {'val': 'F7F8FA'})
                sf.append(sc)
                tcPr.append(sf)

    if notes:
        tb = add_textbox(slide, Inches(0.8), Inches(0.45 * nrows + BODY_TOP.inches + 1.4),
                         Inches(11.7), Inches(1.5))
        tf = tb.text_frame
        for i, note in enumerate(notes):
            if i == 0:
                set_para(tf.paragraphs[0], note, Pt(11), DARK_GRAY)
            else:
                add_para(tf, note, Pt(11), DARK_GRAY)

    add_footer(slide, pres.slides.index(slide) + 1)
    return slide


def code_slide(pres, title, code_text, subtitle=None, notes=None):
    slide = blank_slide(pres)
    title_bar(slide, title, subtitle)

    code_w = Inches(11.7)
    code_h = Inches(3.0)
    code_left = Inches(0.8)
    code_top = BODY_TOP

    shape = add_rect(slide, code_left, code_top, code_w, code_h, CODE_BG)
    shape.shadow.inherit = False

    tb = add_textbox(slide, Inches(1.0), Inches(1.4), Inches(11.3), Inches(2.8))
    tf = tb.text_frame
    tf.word_wrap = True
    set_para(tf.paragraphs[0], code_text, Pt(10), CODE_FG,
             font_name='Courier New', alignment=PP_ALIGN.LEFT)

    if notes:
        note_top = Inches(1.3 + 3.0 + 0.3)
        ntb = add_textbox(slide, Inches(0.8), note_top, Inches(11.7), Inches(2.5))
        ntf = ntb.text_frame
        for i, n in enumerate(notes):
            if i == 0:
                set_para(ntf.paragraphs[0], n, Pt(12), BLACK)
            else:
                add_para(ntf, n, Pt(12), DARK_GRAY)

    add_footer(slide, pres.slides.index(slide) + 1)
    return slide


def mixed_slide(pres, title, elements, subtitle=None):
    """elements: list of ('text', str) or ('bullets', [str])"""
    slide = blank_slide(pres)
    title_bar(slide, title, subtitle)
    tb = add_textbox(slide, Inches(0.8), BODY_TOP, Inches(11.7), Inches(5.5))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for el_type, content in elements:
        if el_type == 'text':
            if first:
                set_para(tf.paragraphs[0], content, Pt(13), BLACK, space_after=Pt(8))
                first = False
            else:
                add_para(tf, content, Pt(13), BLACK, space_after=Pt(8))
        elif el_type == 'heading':
            if first:
                set_para(tf.paragraphs[0], content, Pt(15), BLACK, bold=True,
                         space_after=Pt(4), space_before=Pt(10))
                first = False
            else:
                add_para(tf, content, Pt(15), BLACK, bold=True,
                         space_after=Pt(4), space_before=Pt(10))
        elif el_type == 'bullets':
            for item in content:
                bullet = '\u2022 ' + item
                if first:
                    set_para(tf.paragraphs[0], bullet, Pt(12), DARK_GRAY,
                             space_after=Pt(3))
                    first = False
                else:
                    add_para(tf, bullet, Pt(12), DARK_GRAY, space_after=Pt(3))
    add_footer(slide, pres.slides.index(slide) + 1)
    return slide


def build():
    pres = Presentation()
    pres.slide_width = SLIDE_W
    pres.slide_height = SLIDE_H

    # ══════════════════════════════════════════════════════════════════
    # Slide 1: Cover
    # ══════════════════════════════════════════════════════════════════
    slide = blank_slide(pres)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(4.0), GREEN_DARK)
    tb = add_textbox(slide, Inches(1.5), Inches(1.0), Inches(10.3), Inches(2.5))
    tf = tb.text_frame
    tf.word_wrap = True
    set_para(tf.paragraphs[0], 'GPUDirect RDMA (GDR)\n与 GPUDirect Storage (GDS)\n技术文档',
             Pt(34), WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    tb2 = add_textbox(slide, Inches(1.5), Inches(3.8), Inches(10.3), Inches(1.2))
    tf2 = tb2.text_frame
    set_para(tf2.paragraphs[0], "基于 NVIDIA 官方文档整理", Pt(18),
             RGBColor(0xC8, 0xE6, 0xC9), alignment=PP_ALIGN.CENTER)
    add_para(tf2, "面向工程师与架构师的平衡深度资料", Pt(14),
             RGBColor(0xA5, 0xD6, 0xA7), alignment=PP_ALIGN.CENTER)
    tb3 = add_textbox(slide, Inches(1.5), Inches(5.5), Inches(10.3), Inches(1.0))
    tf3 = tb3.text_frame
    set_para(tf3.paragraphs[0], "GPUDirect RDMA v13.3  |  GPUDirect Storage v1.18", Pt(12),
             DARK_GRAY, alignment=PP_ALIGN.CENTER)
    add_footer(slide, pres.slides.index(slide) + 1)

    # ══════════════════════════════════════════════════════════════════
    # Slide 2: 目录
    # ══════════════════════════════════════════════════════════════════
    toc_items = [
        "一、背景与动机",
        "二、GPUDirect 技术演进",
        "三、软件生态全景",
        "四、控制面 vs 数据面路径",
        "五、GPU 兼容性矩阵",
        "六、GDR 技术细节",
        "七、GDR 全生命周期实例",
        "八、GDS 技术细节",
        "九、GDR vs GDS 对比",
        "十、典型应用场景",
        "十一、部署与启用要点",
        "十二、总结与参考链接",
    ]
    slide = blank_slide(pres)
    title_bar(slide, "目录")
    tb = add_textbox(slide, Inches(1.5), Inches(1.5), Inches(5), Inches(5.5))
    tf = tb.text_frame
    set_para(tf.paragraphs[0], toc_items[0], Pt(15), BLACK, space_after=Pt(8))
    for item in toc_items[1:6]:
        add_para(tf, item, Pt(15), BLACK, space_after=Pt(8))
    tb2 = add_textbox(slide, Inches(7.5), Inches(1.5), Inches(5), Inches(5.5))
    tf2 = tb2.text_frame
    set_para(tf2.paragraphs[0], toc_items[6], Pt(15), BLACK, space_after=Pt(8))
    for item in toc_items[7:]:
        add_para(tf2, item, Pt(15), BLACK, space_after=Pt(8))
    add_footer(slide, pres.slides.index(slide) + 1)

    # ══════════════════════════════════════════════════════════════════
    # Slide 3: 背景与动机 (1) - GPU数据搬运的痛点
    # ══════════════════════════════════════════════════════════════════
    bullet_slide(pres, "一、背景与动机 — GPU 数据搬运的痛点", [
        "GPU 算力增长远快于 I/O 带宽，大模型训练中数据加载成为瓶颈，应用从 compute-bound 退化为 IO-bound",
        "传统路径三重开销：\n"
        "  双倍拷贝 — 数据从存储拷到 CPU bounce buffer（即 CPU 系统内存中的临时中转缓冲区），再拷到 GPU，一次搬运两份数据移动",
        "  带宽瓶颈 — 数据须穿过 CPU 根端口，CPU PCIe 带宽可能低于 GPU/NIC 能力，CPU 成为带宽漏斗",
        "  CPU 占用 — CPU 搬运数据干扰其他计算，引入抖动 (jitter) 与延迟",
    ])

    # ══════════════════════════════════════════════════════════════════
    # Slide 4: 背景与动机 (2) - GPUDirect 家族概览
    # ══════════════════════════════════════════════════════════════════
    table_slide(pres, "一、背景与动机 — GPUDirect 家族概览",
        ["技术", "全称", "直连对象", "说明"],
        [
            ["GDR", "GPUDirect RDMA", "GPU ↔ 网卡 (NIC)",
             "远程系统间 GPU 直接通信，绕过 CPU，10x 性能"],
            ["GDS", "GPUDirect Storage", "GPU ↔ 存储 (NVMe/NVMe-oF)",
             "存储 DMA 引擎直接搬数据进/出 GPU 显存"],
            ["P2P", "GPUDirect Peer-to-Peer", "GPU ↔ GPU",
             "经 PCIe/NVLink 的 GPU 间直接拷贝与 load/store"],
            ["Video", "GPUDirect for Video", "GPU ↔ 采集设备",
             "帧设备与 GPU 间高效传帧（仅 Windows）"],
            ["Async", "GPUDirect Async", "异步触发",
             "存储 I/O 排入 CUDA stream/graph 实现计算-I/O 联动"],
        ],
        col_widths=[Inches(1.0), Inches(2.2), Inches(2.5), Inches(6.0)],
        notes=["本文聚焦 GDR (RDMA) 与 GDS (Storage) 两项核心技术",
               "核心范式：控制面在 CPU 建立通路，数据面由设备 DMA 引擎直连 GPU 显存"]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 5: GPUDirect 技术演进 (1) - 早期
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "二、GPUDirect 技术演进 — 早期探索",
        [
            ("heading", "2010 年前后 (Fermi 时代)"),
            ("text", '首次提出 GPUDirect 概念，允许第三方 PCIe 设备绕过 CPU 直接访问 GPU 显存，奠定"设备直连显存"的技术基础'),
            ("heading", "2011 (CUDA 4.0)"),
            ("text", "引入统一虚拟地址空间 (UVA)，推出 GPUDirect P2P：同一节点内多块 GPU 经 PCIe 直接拷贝与互访"),
            ("heading", "2012 (Kepler + CUDA 5.0)"),
            ("text", "GPUDirect RDMA 正式登场：把直连范围扩展到 GPU 与远端第三方设备之间"),
            ("bullets", [
                '确立"控制面在 CPU，数据面由设备 DMA 直连"的核心范式',
                "明确硬件前提：两端共享同一 PCIe 根复合体、IOMMU 须禁用或直通",
            ]),
        ],
        subtitle="核心演进主线：直连对象不断扩展，实现从专有驱动向 Linux 上游标准靠拢"
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 6: GPUDirect 技术演进 (2) - 发展期
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "二、GPUDirect 技术演进 — 产品化与扩展",
        [
            ("heading", "2014-2016 (CUDA 6.0 - 8.0)"),
            ("bullets", [
                "取消强制安全令牌，降低第三方接入门槛",
                "BAR 映射可在多次锁定间复用，节省资源",
                "补齐 I/O 地址与 CPU 物理地址不一致平台的地址映射能力",
                "GPUDirect Video 将直连对象扩展到帧采集设备",
            ]),
            ("heading", "2018-2021 (CUDA 10.1 - 11.4)"),
            ("bullets", [
                "向嵌入式/车载平台延伸：Jetson AGX Xavier、DRIVE AGX Xavier、Jetson Orin",
                "nvidia-peermem.ko 作为官方 in-tree 模块随驱动发布——标志网络侧直连进入产品化阶段",
                "与 Mellanox OFED、InfiniBand、RoCE 协议栈深度协同",
            ]),
        ],
        subtitle="从实验室到产品化：网络侧直连能力日趋成熟"
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 7: GPUDirect 技术演进 (3) - 现代化
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "二、GPUDirect 技术演进 — 现代化与标准化",
        [
            ("heading", "2020 (CUDA 11.x) — GDS 登场"),
            ("text", "GPUDirect Storage 加入 Magnum IO 体系：存储 DMA 引擎直接搬数据进/出 GPU 显存"),
            ("heading", "2022-2024 (CUDA 12.2-12.8)"),
            ("bullets", [
                '异步 IO 接入 CUDA 流与图，实现"数据就绪后自动启 kernel、kernel 完成后自动存盘"',
                "支持 CPU 与 GPU 间可迁移内存，兼容 Grace-Hopper 统一内存架构",
                "CUDA 12.8 起 NVMe 直连改走 Linux 上游内核 PCI P2PDMA 基础设施",
            ]),
            ("heading", "2025 (CUDA 13.0) — 上游化演进"),
            ("text", "传统 nv-p2p 接口在 Blackwell SoC 进入弃用，NVIDIA 鼓励迁移到 Linux 上游标准接口 (pin_user_pages / dma_map_sgtable)"),
        ],
        subtitle="三条交织脉络：直连对象扩展 → 实现上游化 → 控制面与 CUDA 融合"
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 8: 软件生态 — nvidia.ko
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "三、软件生态全景 — nvidia.ko（GPU 端地址提供者）",
        [
            ("text", "nvidia.ko 是 NVIDIA GPU 内核驱动，是整个 GDR 机制的底层基石。它在操作系统层面的空白地带（Linux 缺乏在驱动之间交换 MMIO 区的标准机制）自行承担地址翻译与映射职责。"),
            ("heading", "核心作用"),
            ("bullets", [
                "页锁定 (pin)：把 GPU 虚拟地址对应显存锁定，确保 DMA 期间不发生迁移或回收",
                "物理地址回填：返回 GPU BAR 总线地址的物理页表数组",
                "DMA 映射：平台 I/O 地址与 CPU 物理地址不一致时进一步映射到设备 I/O 地址空间",
                "撤销通知 (free callback)：GPU 显存释放/进程退出时同步回调第三方驱动",
            ]),
            ("heading", "关键 API（内核态导出符号）"),
            ("text", "nvidia_p2p_get_pages() / nvidia_p2p_put_pages() / nvidia_p2p_dma_map_pages() / nvidia_p2p_free_page_table()"),
        ],
        subtitle='nvidia.ko 本身不感知"是谁在 DMA"——它只负责把地址交出来'
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 9: 软件生态 — nvidia-peermem + MLNX_OFED
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "三、软件生态全景 — nvidia-peermem.ko + MLNX_OFED",
        [
            ("heading", "nvidia-peermem.ko — 胶水模块"),
            ("text", 'CUDA 11.4 起随驱动发布的 in-tree 模块。站在 InfiniBand/RDMA 子系统一侧，把"GPU 显存"伪装成可供 RDMA 注册的"对等内存 (peer memory)"。'),
            ("bullets", [
                "向 IB Core 注册为 peer memory client，回调 nvidia-p2p 获取 GPU BAR 地址",
                "把地址回填给 HCA 驱动 (mlx5_ib)，编程 mkey 转换表",
                "支持 ConnectX-3 VPI 及更新适配器，兼容 InfiniBand 与 RoCE",
                "取代已弃用的社区离线模块 nv_peer_mem",
            ]),
            ("heading", "MLNX_OFED — 核心承载栈"),
            ("text", "绝大多数 GDR 真实部署 (NCCL、MPI) 都跑在 MLNX_OFED 之上。"),
            ("bullets", [
                "提供 RDMA peer-memory 框架与标准 API",
                "ConnectX HCA 驱动接收 GPU BAR 地址后经 PCIe 直接 DMA 读写 GPU 显存",
                "版本协同约束：nvidia-peermem 须用 MLNX_OFED 的 RDMA API 编译",
            ]),
        ],
        subtitle='nvidia-peermem 是 GDR 在 IB 侧的"唯一接入点"'
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 10: 软件生态 — 第三方驱动接入路径
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "三、软件生态全景 — 第三方驱动如何接入 GDR",
        [
            ("heading", "用户态职责"),
            ("bullets", [
                "cudaMalloc 分配 GPU 显存（UVA 空间）",
                "cuPointerGetAttribute() 判定指针归属 (CPU vs GPU)",
                "cuPointerSetAttribute(SYNC_MEMOPS) 启用同步语义",
                "将 GPU VA + 长度经自研驱动 IOCTL 传入内核态",
                "维护 registration cache + lazy unpinning 优化性能",
            ]),
            ("heading", "内核态职责"),
            ("bullets", [
                "编译时链接 nvidia.ko 的 Module.symvers (KBUILD_EXTRA_SYMBOLS)",
                "64KB 对齐后调用 nvidia_p2p_get_pages() 获取物理页表",
                "按需调用 nvidia_p2p_dma_map_pages() 映射到设备 I/O 地址空间",
                "将 scatter-gather 列表写入自研设备 DMA 引擎寄存器",
                "注册 free_callback 处理 GPU 显存提前释放",
            ]),
            ("heading", "接入方式选择"),
            ("text", "Mellanox 网卡 → 复用 nvidia-peermem + MLNX_OFED；非 Mellanox 设备 → 自研内核驱动直接调用 nvidia_p2p_* 符号"),
        ]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 11: 软件生态 — GDS 组件
    # ══════════════════════════════════════════════════════════════════
    table_slide(pres, "三、软件生态全景 — GDS 软件组件",
        ["组件", "来源", "类型", "作用"],
        [
            ["libcufile.so", "NVIDIA", "用户态闭源库", "cuFile API 实现：Driver / 同步 IO / Batch / Stream async"],
            ["libcufile_static.a", "NVIDIA", "静态库", "静态链接场景"],
            ["libcufile_rdma.so", "NVIDIA", "用户态库", "获取 RDMA info (keys, GID, LID) 给 libcufile"],
            ["nvidia-fs.ko", "NVIDIA", "内核驱动 GPLv2", "处理 cuFile IOCTL，DMA callback 翻译 GPU VA"],
            ["libcuda.so", "NVIDIA", "CUDA 驱动", "stream 子集 API 运行时依赖"],
            ["第三方 FS 驱动", "伙伴/厂商", "内核态", "Lustre / WekaFS / NFS / EXAScaler"],
            ["kvikio/RAPIDS/DALI", "开源社区", "库", "C++/Python 高层封装"],
        ],
        col_widths=[Inches(2.2), Inches(1.2), Inches(1.8), Inches(6.5)],
        notes=["nvidia-fs.ko 以 GPLv2 许可发布 — 这是 GDS 内核态集成的关键合规边界",
               "CUDA 12.8 起 NVMe 直连走 upstream PCI P2PDMA，不再需要 nvidia-fs.ko"]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 12: 控制面 vs 数据面 — 三种 Agent + 哲学
    # ══════════════════════════════════════════════════════════════════
    table_slide(pres, "四、控制面 vs 数据面 — 三种数据搬移 Agent",
        ["Agent", "搬移方式", "是否需要 bounce buffer", "说明"],
        [
            ["GPU DMA 引擎", "由 CPU 编程", "需要", "第三方设备通常不暴露内存供他方 DMA"],
            ["CPU load/store", "CPU 在两个设备间拷贝", "需要", "CPU 无法在两个设备间直接拷贝"],
            ["存储侧 DMA 引擎 (NVMe/NIC/RAID)", "GPU BAR 地址暴露给 DMA 引擎", "不需要 (GDS 直连)",
             "NIC 驱动 (Mellanox 等) 支持此路径——GDS 的价值所在"],
        ],
        col_widths=[Inches(2.5), Inches(2.5), Inches(3.0), Inches(3.7)],
        notes=["核心范式：控制面（CPU 侧软件建立通路）与数据面（设备 DMA 引擎直接搬数据）分离",
               "GDS = 让存储侧 DMA 引擎路径在文件系统存储端点可用"]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 13: 控制面 vs 数据面 — GDR 控制流/数据流
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "四、GDR 控制流/数据流分离",
        [
            ("heading", "控制面（CPU 侧：IOCTL + pin/unpin + callback）"),
            ("bullets", [
                "CPU 内存 → 标准 get_user_pages() / put_page()",
                "GPU 内存 → nvidia.ko 提供的 nvidia_p2p_get_pages / put_pages",
                "Unpin Callback：nvidia_p2p_get_pages() 必须提供 callback，NVIDIA 驱动在需撤销时同步调用；callback 内须调用 nvidia_p2p_free_page_table()",
            ]),
            ("heading", "数据面（设备侧：DMA 引擎直接读写 GPU BAR）"),
            ("bullets", [
                "第三方设备 DMA 引擎对 GPU BAR 地址发起 read/write",
                "与访问 system memory 方式相同",
                "NVIDIA 内核驱动仅做地址翻译/映射，DMA 不经 CPU",
            ]),
            ("heading", "死锁风险"),
            ("text", "NVIDIA 驱动持锁后调 callback，第三方也持锁 → 需谨慎设计锁层次"),
        ],
        subtitle="GDR 扩展标准 DMA Transfer：新增用户态 CUDA 库 + NVIDIA 内核驱动"
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 14: GPU 兼容性矩阵
    # ══════════════════════════════════════════════════════════════════
    table_slide(pres, "五、GPU 兼容性矩阵 — GDR vs GDS",
        ["维度", "GDR (RDMA)", "GDS (Storage)"],
        [
            ["架构门槛", "Kepler+ (CUDA 5.0)", "需暴露 BAR1 aperture"],
            ["产品线 (能用)", "Tesla (数据中心) / Quadro",
             "Quadro / RTX (专业卡) / 数据中心 GPU (暴露 BAR1 子集)"],
            ["不支持", "GeForce 消费卡；Blackwell SoC (nv-p2p 已移除)", "GeForce 消费卡；不暴露 BAR1 的 GPU"],
            ["BAR1", "Kepler 最小 256MB；Tesla large BAR 16GB+",
             "传输超 aperture 自动分块+GPU 内 staging"],
            ["检测属性", "CU_DEVICE_ATTRIBUTE_GPU_DIRECT_RDMA_SUPPORTED",
             "CU_DEVICE_ATTRIBUTE_DMA_BUF_SUPPORTED"],
            ["内存要求", "cudaMalloc (pinned)；Unified Memory 不显式支持",
             "cudaMalloc / cuMemAlloc / cuMemCreate；cudaMallocManaged 走兼容路径"],
        ],
        col_widths=[Inches(2.0), Inches(5.7), Inches(4.0)],
        notes=["GDS 需要 GPU 暴露 BAR1 aperture：仅 NVIDIA RTX (专业卡) 与数据中心 GPU 的子集满足",
               "Blackwell SoC：CUDA 13.0 起 nv-p2p 不再支持 → 迁移到 Linux 上游内核 API"]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 15: GDR 全生命周期 — 跨节点示意图 + 阶段0
    # ══════════════════════════════════════════════════════════════════
    code_slide(pres, "七、GDR 全生命周期 — 跨节点数据面示意",
        textwrap.dedent("""\
        节点 A                              节点 B
        +-----------+                      +-----------+
        | 用户应用    |   控制面：            | 用户应用    |
        | ibv_post   |   rkey+VA 互换       | (被动方)    |
        | _send(WR)  | <==============>    |            |
        +-----+-----+                      +-----+-----+
              |                                    |
        ------+-- 内核 ---------------------------+-- 内核 --
              |                                    |
        +-----+-----+                      +-----+-----+
        | nvidia.ko  |  pin GPU 页          | nvidia.ko  |  pin GPU 页
        |+peermem.ko | <-- 控制面           |+peermem.ko | <-- 控制面
        +-----+-----+                      +-----+-----+
              |                                    |
        ------+-- 数据面 -------------------------+-- 数据面 --
              |         +-----------+             |
        +-----+---+ DMA | ConnectX  | 网络   +---+-----+
        | GPU A    |<--->| HCA (A端) |------->| GPU B    |
        | BAR1     | 读源 +-----------+ 写目的 | BAR1     |
        +----------+                          +---------+
        数据面全程不经任何一侧 CPU 内存"""),
        subtitle="RDMA WRITE：节点 A 写节点 B 的 GPU",
        notes=["阶段 0 | 启动与就绪 — 搭好前置条件",
               "BIOS 分配 BAR 窗口 → OS 枚举 PCIe 拓扑、配置 IOMMU (须禁用或直通)",
               "加载 nvidia.ko → MLNX_OFED 栈 (ib_core + mlx5_ib) → modprobe nvidia-peermem",
               "nvidia-peermem 向 ib_core 注册为 GPU peer memory client"]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 16: GDR 全生命周期 — 阶段 1+2
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "七、GDR 全生命周期 — 阶段 1 内存注册 + 阶段 2 元数据交换",
        [
            ("heading", "阶段 1 | 本地内存注册 (pin + mkey 编程)"),
            ("text", "目标：GPU VA → GPU 物理/BAR 地址 → HCA mkey 转换表，产出 lkey/rkey"),
            ("bullets", [
                "cudaMalloc 分配 GPU 显存，返回 UVA 空间中的 GPU 虚拟地址",
                "cuPointerSetAttribute(SYNC_MEMOPS) 建立 CUDA 侧内存序一致性契约",
                "ibv_alloc_pd / ibv_reg_mr 注册 Memory Region",
                "ib_core 遍历 peer client 链表，nvidia-peermem 认领 GPU 内存",
                "nvidia_p2p_get_pages() pin 住显存、返回 BAR 总线地址页表",
                "mlx5_ib 编程 HCA 硬件 mkey 转换表 → 产出 lkey / rkey",
            ]),
            ("heading", "阶段 2 | 跨节点元数据交换"),
            ("text", '纯控制面：两端互换 (rkey, 注册 VA)——RDMA WRITE/READ 寻址远端的"门牌号"'),
            ("bullets", [
                "创建 RC QP (Reliable Connection Queue Pair)",
                "经 send/recv 互换 rkey + Remote VA",
                "本步不涉及 nvidia.ko、nvidia-peermem、HCA mkey 修改",
            ]),
        ]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 17: GDR 全生命周期 — 阶段 3+4
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "七、GDR 全生命周期 — 阶段 3 数据传输 + 阶段 4 销毁回收",
        [
            ("heading", "阶段 3 | RDMA WRITE — 节点 A 写节点 B 的 GPU"),
            ("bullets", [
                "A 端 ibv_post_send(WR=RDMA_WRITE) 异步提交，sge={VA, len, lkey}, remote={rkey, VA}",
                "A 端 HCA DMA 引擎用 lkey 查 mkey 表 → PCIe 读 A 的 GPU BAR → 网络发送",
                "B 端 HCA 拆包 → 查自己的 mkey 表 → PCIe 直写 B 的 GPU BAR",
                "GPU BAR ←PCIe DMA→ HCA ←网络→ HCA ←PCIe DMA→ GPU BAR — 零 CPU 拷贝",
            ]),
            ("heading", "RDMA READ 方向镜像：A 端主动拉取 B 的 GPU 数据，B 端 CPU 同样不参与"),
            ("heading", "阶段 4 | 销毁与回收 — 安全拆掉映射链"),
            ("bullets", [
                "正常路径：dereg_mr → nvidia_p2p_put_pages() → cudaFree",
                "异常路径 (cudaFree 先于 dereg_mr)：free_callback 同步触发 → 等 DMA 完成 → nvidia_p2p_free_page_table()",
                "进程退出时 callback 兜底，模块卸载反向 nvidia-peermem → MLNX_OFED → nvidia.ko",
            ]),
        ]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 18: GDR vs GDS 总览对比
    # ══════════════════════════════════════════════════════════════════
    table_slide(pres, "九、GDR vs GDS 总览对比",
        ["维度", "GDR (RDMA)", "GDS (Storage)"],
        [
            ["目标设备", "网卡 (NIC) / 采集卡 / 存储适配器", "NVMe / NVMe-oF / NIC / RAID"],
            ["直连对象", "远程系统间 GPU", "GPU ↔ 存储"],
            ["用户态 API", "cuPointerSetAttribute + 第三方通信库", "cuFile API (cuFileRead/Write 等)"],
            ["内核模块", "nvidia.ko + nvidia-peermem", "nvidia-fs.ko (GPLv2)"],
            ["同步语义", "SYNC_MEMOPS + CUDA 同步 API", "CUDA stream 语义 / 同步调用"],
            ["典型场景", "多 GPU 集群通信、NCCL、MPI", "数据加载、checkpoint、存储密集型"],
            ["拓扑要求", "同一 PCIe root complex; IOMMU 直通", "同一 PCIe switch 最优；跨 root port 走动态路由"],
            ["依赖", "MLNX_OFED、ConnectX-3+", "kernel 4.15+、MLNX_OFED 5.1+、ConnectX-5+"],
        ],
        col_widths=[Inches(1.8), Inches(5.3), Inches(4.6)],
        notes=["共同点：控制面在 CPU 建通路，数据面由设备 DMA 直连 GPU 显存，绕过 CPU bounce buffer",
               "GDR 面向网络通信，GDS 面向存储 IO"]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 19: 典型应用场景
    # ══════════════════════════════════════════════════════════════════
    table_slide(pres, "十、典型应用场景",
        ["场景", "价值"],
        [
            ["AI 训练数据加载", "大数据集从存储直达 GPU；DALI/kvikio 加速数据流水线"],
            ["推理服务", "模型权重/KV-cache 从存储直达 GPU，降低首 token 延迟"],
            ["HPC 模拟 checkpoint", "大规模 GPU 状态直接写存储，绕过 CPU，缩短 checkpoint 窗口"],
            ["多 GPU 集群通信 (GDR)", "NCCL/RDMA/MPI 分布式训练，GPU 间绕过 CPU 直接通信，10x 性能"],
            ["推荐系统 (DLRM)", "存储密集型场景，IO 直接进 GPU，消除 CPU 中转瓶颈"],
        ],
        col_widths=[Inches(3.0), Inches(8.7)],
        notes=["适用条件 (GDS)：数据传输直接到/出 GPU；IO 显著瓶颈；使用 cudaMalloc + cuFile API",
               "不适用：应用在 GPU 计算前后用 CPU 解析数据；计算时间远大于 IO 时间"]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 20: 部署与启用要点
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "十一、部署与启用要点",
        [
            ("heading", "GDR 部署"),
            ("bullets", [
                "安装 MLNX_OFED (带 RDMA peer memory 支持) + NVIDIA GPU 驱动",
                "modprobe nvidia-peermem (注意：无自动加载服务，需手动)",
                "检查拓扑：lspci -t；IOMMU 须禁用或 pass-through",
                "第三方模块编译时链接 nvidia.ko 的 Module.symvers",
            ]),
            ("heading", "GDS 部署"),
            ("bullets", [
                "安装 CUDA Toolkit (含 libcufile.so) + nvidia-fs 内核驱动 (DKMS)",
                "依赖：libmount-dev、libnuma-dev；RDMA FS：MLNX_OFED 5.1+, ConnectX-5+",
                "配置 cufile.json (动态路由、mount_table、log level)",
                "验证：gdscheck -p；监控 BAR1：nvidia-smi -q",
            ]),
            ("heading", "已知坑"),
            ("bullets", [
                "CUDA 12.2 race bug: R515-R535 invalidation callback 与 put_pages 并发 → 升级驱动",
                "Blackwell SoC: nv-p2p 不可用 → 迁移 pin_user_pages / dma_map_sgtable",
                "容器化：宿主机内核态组件 (nvidia-fs.ko/nvidia-peermem)，缺驱动时启用兼容模式",
            ]),
        ]
    )

    # ══════════════════════════════════════════════════════════════════
    # Slide 21: 总结
    # ══════════════════════════════════════════════════════════════════
    mixed_slide(pres, "十二、总结 — 核心要点",
        [
            ("text", "GDR 与 GDS 同属 GPUDirect 家族，核心范式一致：控制面在 CPU 建通路，数据面由设备 DMA 引擎直连 GPU 显存，绕过 CPU bounce buffer。"),
            ("heading", "关键区别"),
            ("bullets", [
                "GDR 面向网络 (NIC)，用 nvidia-peermem + nvidia_p2p_* API",
                "GDS 面向存储，用 cuFile API + nvidia-fs.ko (GPLv2)",
            ]),
            ("heading", "GPU 兼容性"),
            ("bullets", [
                "GDR: Tesla/Quadro (Kepler+)",
                "GDS: 暴露 BAR1 的 RTX/数据中心 GPU 子集",
                "GeForce 消费卡均不支持",
            ]),
            ("heading", "演进方向"),
            ("text", "从专有 nv-p2p 迁移到 Linux 上游 pin_user_pages / dma_map_sgtable；NVMe 路径走 upstream PCI P2PDMA"),
            ("heading", "参考链接"),
            ("text", "GPUDirect 主页 | GPUDirect RDMA 文档 | GDS Overview / Design Guide | Magnum IO | nv_peer_memory (GitHub)"),
        ]
    )

    # ── figure slides ─────────────────────────────────────────────────
    fig_titles = {
        'fig1.svg': '图 1：传统路径 vs GPUDirect 直连路径对比',
        'fig2.svg': '图 2：GDR 软件栈分层架构',
        'fig3.svg': '图 3：第三方驱动接入 GDR 的数据路径',
        'fig4.svg': '图 4：跨节点 GDR 数据面（RDMA WRITE）时序流程',
        'fig5.svg': '图 5：GDS 软件栈分层架构与 Workflow 1 调用栈',
        'fig6.svg': '图 6：Phoenix 架构对比',
        'fig7.svg': '图 7：GDR 全生命周期五阶段时序',
        'fig8.svg': '图 8：控制面 vs 数据面核心范式',
    }
    for fname, title in fig_titles.items():
        svg_path = os.path.join(FIGS, fname)
        if not os.path.exists(svg_path):
            continue
        try:
            png_data = cairosvg.svg2png(url=svg_path)
        except Exception:
            continue
        slide = pres.slides.add_slide(pres.slide_layouts[6])  # blank
        title_bar(slide, title)
        # Position image centered on slide
        png_stream = io.BytesIO(png_data)
        slide.shapes.add_picture(png_stream, Inches(0.5), Inches(1.0),
                                 Inches(12.3), Inches(6.0))

    # ── finalize ─────────────────────────────────────────────────────
    pres.save(DST)
    print(f"Built {DST} ({len(pres.slides)} slides)")


if __name__ == '__main__':
    build()
